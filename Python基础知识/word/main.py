# coding:utf-8
from docx import Document

document = Document()
tables = document.add_table(rows=2, cols=6, style = 'Table Grid')
# 创建表格，rows是行数，cols是列数

tables.cell(0,1).text = (u'土拨鼠')
# 第0行的第1列（就是第一行的第二列，计算机0开始计数）
# 的内容是 土拨鼠
tables.cell(1,1).text=(u'啊啊啊啊啊')
# 第1行的第1列（就是第二行的第二列）
# 的内容是啊啊啊啊啊
document.save('test2.docx')